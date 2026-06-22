"""PQE Wave A — _write_xlsx, _round_trip_validate, and _value_from_snapshot.

Covers:
  PQE-BE-32  _write_xlsx header + value rows (real openpyxl round-trip)
  PQE-BE-33  _write_xlsx sheet-title sanitization (>31 chars / colon → dash)
  PQE-BE-34  _round_trip_validate pass (freshly written bytes + original row)
  PQE-BE-35  _round_trip_validate header mismatch → passed=False + diagnostic
  PQE-BE-36  _round_trip_validate value mismatch → passed=False + mismatches list
  PQE-BE-42  _value_from_snapshot precedence (fields_jsonb > ai.value > "")

All tests are ``unit`` — no DB, no real GCS, no network.
Uses real openpyxl (it is an offline, deterministic dependency — see plan §5.4).
"""

from __future__ import annotations

from io import BytesIO

import pytest

pytestmark = pytest.mark.unit

pytest.importorskip("openpyxl", reason="openpyxl is required for XLSX round-trip tests")

import openpyxl  # noqa: E402 — import after importorskip guard

from app.modules.export.domain import XlsxColumnSpec, XlsxRowSpec  # noqa: E402
from app.modules.export.service import (  # noqa: E402
    _round_trip_validate,
    _value_from_snapshot,
    _write_xlsx,
)


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────
def _make_row(
    *,
    label: str = "Test Sheet",
    headers: list[str],
    values: list,
) -> XlsxRowSpec:
    """Build a minimal XlsxRowSpec for the given headers/values."""
    cols = tuple(
        XlsxColumnSpec(
            canonical_name=f"col_{i}",
            meesho_column_header=h,
            meesho_column_index=i,
            value=v,
        )
        for i, (h, v) in enumerate(zip(headers, values, strict=True))
    )
    return XlsxRowSpec(main_sheet_label=label, columns=cols)


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-32  _write_xlsx header + value rows
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_32_write_xlsx_has_two_rows():
    """_write_xlsx produces a re-loadable XLSX with exactly 2 rows (header + value).

    Arrange: XlsxRowSpec with 3 columns.
    Act: _write_xlsx(row).
    Assert: re-opened workbook has 2 rows; row 1 = headers; row 2 = values.
    """
    row = _make_row(
        headers=["Product Name", "Colour", "Size"],
        values=["Test Kurti", "Red", "M"],
    )

    xlsx_bytes = _write_xlsx(row)

    assert isinstance(xlsx_bytes, bytes) and len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))

    assert len(all_rows) == 2, (
        f"Expected exactly 2 rows (header + value), got {len(all_rows)}"
    )
    assert list(all_rows[0]) == ["Product Name", "Colour", "Size"], (
        f"Header row mismatch: {all_rows[0]!r}"
    )
    assert list(all_rows[1]) == ["Test Kurti", "Red", "M"], (
        f"Value row mismatch: {all_rows[1]!r}"
    )


def test_pqe_be_32_write_xlsx_headers_match_meesho_column_headers():
    """_write_xlsx uses meesho_column_header (not canonical_name) for XLSX headers.

    The Meesho template preserves typo headers (e.g. 'Clour' not 'Colour')
    via the seed pipeline — the XLSX must emit the header from meesho_column_header.
    """
    row = _make_row(
        headers=["Clour"],  # intentional typo-preserve header
        values=["Red"],
    )

    xlsx_bytes = _write_xlsx(row)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    first_row = next(ws.iter_rows(values_only=True))

    assert first_row[0] == "Clour", (
        f"XLSX must emit the meesho_column_header verbatim; got {first_row[0]!r}"
    )


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-33  _write_xlsx sheet-title sanitization
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_33_write_xlsx_title_truncated_at_31_chars():
    """Sheet title longer than 31 chars is truncated to 31 (openpyxl constraint).

    Arrange: main_sheet_label = 32-char string.
    Act: _write_xlsx.
    Assert: worksheet title is exactly 31 chars.
    """
    long_label = "A" * 32  # 32 chars — exceeds openpyxl 31-char limit
    row = _make_row(label=long_label, headers=["Col"], values=["Val"])

    xlsx_bytes = _write_xlsx(row)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)

    assert len(wb.active.title) == 31, (
        f"Sheet title must be truncated to 31 chars; got {len(wb.active.title)!r}"
    )


def test_pqe_be_33_write_xlsx_colon_replaced_with_dash():
    """Colons in sheet title are replaced with dashes (openpyxl sheet name constraint).

    Arrange: main_sheet_label contains a colon.
    Act: _write_xlsx.
    Assert: no colon in the worksheet title.
    """
    row = _make_row(
        label="Cat: Sarees & Suits",
        headers=["Name"],
        values=["Kurti"],
    )

    xlsx_bytes = _write_xlsx(row)
    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)

    assert ":" not in wb.active.title, (
        f"Sheet title must not contain colons; got {wb.active.title!r}"
    )
    assert "-" in wb.active.title, (
        f"Colon should be replaced with dash; got {wb.active.title!r}"
    )


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-34  _round_trip_validate pass
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_34_round_trip_validate_passes_on_good_bytes():
    """_round_trip_validate(bytes, row) returns passed=True on freshly written XLSX.

    This is the primary guarantee: the bytes _write_xlsx produces re-parse to
    the exact same headers and values that went in.
    """
    row = _make_row(
        headers=["Product Name", "Price"],
        values=["Test Saree", "499"],
    )
    xlsx_bytes = _write_xlsx(row)

    result = _round_trip_validate(xlsx_bytes, row)

    assert result.passed is True, (
        f"Round-trip must pass on freshly written bytes; diagnostic: {result.diagnostic!r}"
    )
    assert result.mismatches == (), (
        f"No mismatches expected on a clean round-trip; got {result.mismatches!r}"
    )


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-35  _round_trip_validate header mismatch → passed=False + diagnostic
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_35_round_trip_validate_fails_on_header_mismatch():
    """Tampered XLSX (wrong header) causes round-trip to fail with diagnostic.

    Arrange: write a valid row; then tamper the header in a new workbook.
    Act: _round_trip_validate with the tampered bytes and the original row.
    Assert: passed=False; diagnostic mentions the header mismatch.
    """
    row = _make_row(
        headers=["Product Name", "Colour"],
        values=["Kurti", "Red"],
    )

    # Produce tampered bytes: swap header names.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["WRONG HEADER", "Colour"])  # first header tampered
    ws.append(["Kurti", "Red"])
    buf = BytesIO()
    wb.save(buf)
    tampered_bytes = buf.getvalue()

    result = _round_trip_validate(tampered_bytes, row)

    assert result.passed is False, "Header mismatch must cause round-trip failure"
    assert result.diagnostic is not None and len(result.diagnostic) > 0, (
        "A header-mismatch diagnostic message must be present"
    )


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-36  _round_trip_validate value mismatch → passed=False + mismatches
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_36_round_trip_validate_fails_on_value_mismatch():
    """Tampered XLSX (wrong value) causes round-trip to fail, listing canonical names.

    Arrange: write a valid row; tamper the value row in a new workbook.
    Act: _round_trip_validate with tampered bytes and original row.
    Assert: passed=False; mismatches is non-empty.
    """
    row = _make_row(
        headers=["Product Name", "Colour"],
        values=["Kurti", "Red"],
    )

    # Produce tampered bytes: value in col 0 is wrong.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["Product Name", "Colour"])  # correct headers
    ws.append(["WRONG VALUE", "Red"])  # first value tampered
    buf = BytesIO()
    wb.save(buf)
    tampered_bytes = buf.getvalue()

    result = _round_trip_validate(tampered_bytes, row)

    assert result.passed is False, "Value mismatch must cause round-trip failure"
    assert len(result.mismatches) > 0, (
        "mismatches must list the canonical names of tampered columns"
    )
    assert "col_0" in result.mismatches, (
        f"'col_0' (tampered column) must appear in mismatches; got {result.mismatches!r}"
    )


def test_pqe_be_36_round_trip_validate_fails_on_fewer_than_2_rows():
    """XLSX with only 1 row (header only, no value row) → passed=False.

    This guards against empty or truncated exports.
    """
    row = _make_row(
        headers=["Product Name"],
        values=["Kurti"],
    )

    # Build a 1-row XLSX (header only).
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["Product Name"])  # header row only
    buf = BytesIO()
    wb.save(buf)
    one_row_bytes = buf.getvalue()

    result = _round_trip_validate(one_row_bytes, row)

    assert result.passed is False, "<2-row XLSX must fail round-trip validation"
    assert result.diagnostic is not None


# ─────────────────────────────────────────────────────────────────────────────
# PQE-BE-42  _value_from_snapshot precedence
# ─────────────────────────────────────────────────────────────────────────────
def test_pqe_be_42_fields_jsonb_takes_precedence_over_ai():
    """fields_jsonb value takes precedence over ai_suggestions.

    Arrange: snapshot with both fields and ai_suggestions for the same canonical.
    Act: _value_from_snapshot("product_name", snapshot).
    Assert: returns the fields_jsonb value, not the ai suggestion.
    """
    from types import SimpleNamespace

    snapshot = SimpleNamespace(
        fields={"product_name": "Fields Value"},
        ai_suggestions={
            "product_name": {"value": "AI Suggestion"},
        },
    )

    result = _value_from_snapshot("product_name", snapshot)

    assert result == "Fields Value", (
        f"fields_jsonb must take precedence over ai_suggestions; got {result!r}"
    )


def test_pqe_be_42_ai_suggestions_value_used_when_fields_missing():
    """When fields_jsonb lacks a canonical, ai_suggestions[canonical].value is used.

    Arrange: snapshot with no fields entry but an ai_suggestions dict entry.
    Act: _value_from_snapshot("description", snapshot).
    Assert: returns the .value from ai_suggestions.
    """
    from types import SimpleNamespace

    snapshot = SimpleNamespace(
        fields={},
        ai_suggestions={
            "description": {"value": "AI-generated description"},
        },
    )

    result = _value_from_snapshot("description", snapshot)

    assert result == "AI-generated description", (
        f"ai_suggestions[canonical].value should be used when fields missing; got {result!r}"
    )


def test_pqe_be_42_missing_entirely_returns_empty_string():
    """When the canonical is absent from both fields and ai_suggestions, returns "".

    The XLSX cell renders blank per §14.K fixture 12.
    """
    from types import SimpleNamespace

    snapshot = SimpleNamespace(fields={}, ai_suggestions={})

    result = _value_from_snapshot("nonexistent_canonical", snapshot)

    assert result == "", (
        f"Missing canonical must return '' for blank XLSX cell; got {result!r}"
    )


def test_pqe_be_42_none_field_falls_back_to_ai():
    """A None value in fields_jsonb is treated as absent; ai_suggestions used.

    Ensure that a None stored under a canonical in fields_jsonb is not returned
    as None (which would appear as "None" in the XLSX cell — wrong behaviour).
    """
    from types import SimpleNamespace

    snapshot = SimpleNamespace(
        fields={"color": None},
        ai_suggestions={"color": {"value": "Red"}},
    )

    result = _value_from_snapshot("color", snapshot)

    assert result == "Red", (
        f"None in fields_jsonb must fall through to ai_suggestions; got {result!r}"
    )
